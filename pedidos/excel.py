import io
import zipfile
from decimal import Decimal

from django.conf import settings
from django.http import HttpResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import IllegalCharacterError

TEMPLATE_PATH = settings.BASE_DIR / 'plantillas' / 'formato_nota_pedido.xlsx'


def _set_value(ws, cell, value):
    """Escribe `value` en `cell` limpiando caracteres no soportados por Excel."""
    if value is None or value == '':
        ws[cell] = None
        return
    if isinstance(value, str):
        value = value.replace('\x00', '')
    try:
        ws[cell] = value
    except IllegalCharacterError:
        ws[cell] = ''.join(c for c in str(value) if ord(c) >= 32)


# Formato numérico de Excel para fechas. Se escribe la fecha como valor real
# (no texto) para que Excel la ordene y sume bien, y el formato decide cómo se
# ve: dd/mm/aaaa.
EXCEL_DATE_FORMAT = 'dd/mm/yyyy'


def _fmt_fecha(fecha):
    """Fecha para el Excel.

    Se devuelve el objeto ``date``/``datetime`` en vez de un string ISO: como
    texto, Excel lo muestra tal cual y no lo reconoce como fecha (no se puede
    ordenar ni usar en fórmulas). El formato visible lo aplica
    ``_set_fecha``.
    """
    return fecha or None


def _set_fecha(ws, cell, fecha):
    """Escribe una fecha real en ``cell`` con formato dd/mm/aaaa."""
    fecha = _fmt_fecha(fecha)
    if fecha is None:
        ws[cell] = None
        return
    _set_value(ws, cell, fecha)
    ws[cell].number_format = EXCEL_DATE_FORMAT


def build_nota_workbook(nota):
    wb = load_workbook(TEMPLATE_PATH)
    ws = wb.active

    # ---- Cabecera / datos generales --------------------------------------
    _set_value(ws, 'D13', nota.obra.nombre if nota.obra else '')
    _set_value(ws, 'K13', nota.numero)
    _set_value(ws, 'D14', nota.encargado)
    _set_value(ws, 'H14', nota.etapa)
    _set_fecha(ws, 'K14', nota.fecha)
    _set_value(ws, 'D15', nota.solicitante)
    _set_value(ws, 'H15', nota.prioridad)
    _set_fecha(ws, 'K15', nota.limite)
    _set_value(ws, 'D18', nota.proveedor.nombre if nota.proveedor else '')
    _set_value(ws, 'H18', nota.vendedor)

    # ---- Tabla de mercaderías (filas 21 a 31) -----------------------------
    row = 21
    for item in nota.items.all():
        if row > 31:
            break
        _set_value(ws, f'C{row}', int(item.stock) if item.stock not in (None, '') else '')
        _set_value(ws, f'D{row}', float(item.cantidad) if item.cantidad not in (None, '') else '')
        _set_value(ws, f'E{row}', item.detalle)
        _set_value(ws, f'H{row}', item.especificaciones)
        _set_value(ws, f'J{row}', item.estado)
        _set_value(ws, f'K{row}', float(item.precio) if item.precio not in (None, '') else '')
        row += 1

    # ---- Firmas y pie -----------------------------------------------------
    _set_value(ws, 'C33', nota.solicitante)
    if nota.estado == 'aprobada' and nota.aprobado_por:
        _set_value(ws, 'H33', nota.aprobado_por)
    _set_value(ws, 'C34', 'OBSERVACIONES: ' + nota.observaciones)
    _set_value(ws, 'F34', 'Sí' if nota.para_envio else '')
    _set_value(ws, 'F35', 'Sí' if nota.a_retirar else '')
    _set_value(ws, 'H35', nota.retira)

    return wb


def nota_to_bytes(nota):
    wb = build_nota_workbook(nota)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def excel_response(nota):
    data = nota_to_bytes(nota)
    nombre = nota.obra.nombre if nota.obra else ''
    filename = f'NP_{nota.numero:03d}_{nombre}.xlsx'
    response = HttpResponse(
        data,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


def zip_response(notas):
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, 'w', zipfile.ZIP_DEFLATED) as zf:
        for nota in notas:
            nombre = nota.obra.nombre if nota.obra else ''
            filename = f'NP_{nota.numero:03d}_{nombre}.xlsx'
            zf.writestr(filename, nota_to_bytes(nota))
    response = HttpResponse(buffer.getvalue(), content_type='application/zip')
    response['Content-Disposition'] = 'attachment; filename="notas_pedido.zip"'
    return response


# ---- Orden de compra al proveedor -----------------------------------------
# Esta hoja se arma desde cero (a diferencia de la nota, que usa la plantilla
# oficial): el pedido al proveedor no tiene un formato fijo impuesto y lleva
# datos que la nota no tiene (RUC del proveedor, totales, notas de origen).

def _estilo_encabezado(ws, fila, columnas):
    for indice in columnas:
        celda = ws.cell(row=fila, column=indice)
        celda.font = Font(bold=True, size=12)
        celda.alignment = Alignment(horizontal='center')


def orden_compra_workbook(orden):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Orden de Compra'

    empresa = settings.COMPANY
    proveedor = orden.proveedor

    # ---- Encabezado ------------------------------------------------------
    ws['A1'] = empresa['nombre']
    ws['A1'].font = Font(bold=True, size=16)
    ws.merge_cells('A1:G1')
    ws['A1'].alignment = Alignment(horizontal='center')
    ws['A2'] = 'ORDEN DE COMPRA'
    ws['A2'].font = Font(bold=True, size=13)
    ws.merge_cells('A2:G2')
    ws['A2'].alignment = Alignment(horizontal='center')

    ws['A4'] = 'Proveedor:'
    ws['B4'] = proveedor.nombre
    ws['A5'] = 'RUC:'
    ws['B5'] = proveedor.ruc or ''
    ws['A6'] = 'Teléfono:'
    ws['B6'] = proveedor.telefono or ''
    ws['A7'] = 'Dirección:'
    ws['B7'] = proveedor.direccion or ''
    ws['A4'].font = Font(bold=True)
    ws['A5'].font = Font(bold=True)
    ws['A6'].font = Font(bold=True)
    ws['A7'].font = Font(bold=True)

    ws['D4'] = 'O.C. N°'
    ws['E4'] = f'{orden.numero:04d}'
    ws['D5'] = 'Fecha'
    _set_fecha(ws, 'E5', orden.fecha)
    ws['D6'] = 'Emitido por'
    ws['E6'] = (
        orden.creado_por.get_full_name() or orden.creado_por.get_username()
        if orden.creado_por_id else ''
    )
    for etiqueta in ('D4', 'D5', 'D6'):
        ws[etiqueta].font = Font(bold=True)

    # ---- Tabla de materiales --------------------------------------------
    fila = 9
    encabezados = [
        'Material', 'Especificación', 'Unidad', 'Cantidad',
        'P. Unitario', 'Subtotal', 'Observaciones',
    ]
    for indice, titulo in enumerate(encabezados, start=1):
        ws.cell(row=fila, column=indice, value=titulo)
    _estilo_encabezado(ws, fila, range(1, len(encabezados) + 1))

    fila += 1
    notas_numeros = ', '.join(
        f'NP {n.numero}' for n in orden.notas.order_by('numero')
    )
    total_general = Decimal('0')
    for linea in orden.lineas.all():
        ws.cell(row=fila, column=1, value=linea.material.nombre)
        ws.cell(row=fila, column=2, value=linea.especificacion)
        ws.cell(row=fila, column=3, value=linea.material.unidad)
        celda = ws.cell(row=fila, column=4)
        celda.value = float(linea.cantidad)
        celda.number_format = '#,##0.00'
        precio = linea.precio if linea.precio is not None else Decimal('0')
        celda_precio = ws.cell(row=fila, column=5, value=float(precio))
        celda_precio.number_format = '#,##0.00'
        subtotal = linea.cantidad * precio
        total_general += subtotal
        celda_subtotal = ws.cell(row=fila, column=6, value=float(subtotal))
        celda_subtotal.number_format = '#,##0.00'
        fila += 1

    if orden.lineas.exists():
        ws.cell(row=fila, column=5, value='TOTAL').font = Font(bold=True)
        celda_total = ws.cell(row=fila, column=6, value=float(total_general))
        celda_total.font = Font(bold=True)
        celda_total.number_format = '#,##0.00'
        ws.cell(row=fila, column=7, value=orden.observaciones or '')
        fila += 1
    elif orden.observaciones:
        ws.cell(row=fila, column=1, value='Observaciones:').font = Font(bold=True)
        ws.cell(row=fila, column=7, value=orden.observaciones)
        fila += 1

    for indice, ancho in enumerate([30, 20, 10, 14, 14, 16, 34], start=1):
        ws.column_dimensions[get_column_letter(indice)].width = ancho

    # ---- Pie -------------------------------------------------------------
    if notas_numeros:
        fila += 2
        ws.cell(row=fila, column=1, value='Notas de pedido de origen:').font = Font(bold=True)
        ws.cell(row=fila + 1, column=1, value=notas_numeros)

    fila += 3
    ws.cell(row=fila, column=3, value='Firma proveedor').alignment = Alignment(
        horizontal='center'
    )
    ws.cell(row=fila + 3, column=3, value='Recibido conforme').alignment = Alignment(
        horizontal='center'
    )
    for offset in (1, 2):
        ws.cell(row=fila + offset, column=3).border = Border(top=Side(style='thin'))
    ws.cell(row=fila + 3, column=3).border = Border(top=Side(style='thin'))

    return wb


def orden_compra_response(orden):
    buffer = io.BytesIO()
    orden_compra_workbook(orden).save(buffer)
    nombre_proveedor = ''.join(
        c if c.isalnum() else '-' for c in orden.proveedor.nombre
    ).strip('-')
    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = (
        f'attachment; filename="OC_{orden.numero:04d}_{nombre_proveedor}.xlsx"'
    )
    return response