"""Consolidado de compras y órdenes al proveedor."""

from datetime import date
from decimal import Decimal

from django.urls import reverse

from pedidos.models import Material, NotaPedido, Obra, OrdenCompra, Proveedor

from .base import BasePermisosTest, item, linea, material, orden


class BaseComprasTest(BasePermisosTest):
    def setUp(self):
        self.client.force_login(self.admin)
        self.cemento = material('Cemento', unidad='bolsa')
        self.ladrillo = material('Ladrillo hueco', unidad='unidad')
        # nota_ana y nota_beto en obras distintas: el consolidado tiene que
        # desglosar por ambas.
        self.obra_beto = Obra.objects.create(nombre='Obra Norte')
        self.nota_beto.obra = self.obra_beto
        self.nota_beto.save(update_fields=['obra'])
        item(self.nota_ana, 'Cemento', 30)
        item(self.nota_ana, 'Ladrillo hueco', 100)
        item(self.nota_beto, 'Cemento', 15)
        self.proveedor = Proveedor.objects.first()

    def fila_de(self, respuesta, material_obj):
        for total, _obras in respuesta.context['totales']:
            if total['material_id'] == material_obj.pk:
                return total
        return None

    def crear_orden(self, materiales, **extra):
        datos = {
            'materiales': [str(m.pk) for m in materiales],
            'orden_proveedor': str(self.proveedor.pk),
            'orden_fecha': '2026-08-05',
            'orden_observaciones': 'Entrega en 5 días',
        }
        datos.update(extra)
        return self.client.post(reverse('pedidos:compras_crear_orden'), datos)

    def _post_orden(self, **datos):
        base = {
            'orden_proveedor': str(self.proveedor.pk),
            'orden_fecha': '2026-08-05',
            'orden_observaciones': '',
        }
        base.update(datos)
        return self.client.post(reverse('pedidos:compras_crear_orden'), base)


class ConsolidadoTest(BaseComprasTest):
    def test_agrupa_por_material_sumando_cantidades(self):
        respuesta = self.client.get(reverse('pedidos:compras'))
        cemento = self.fila_de(respuesta, self.cemento)
        self.assertEqual(cemento['total'], Decimal('45.00'))
        self.assertEqual(cemento['notas'], 2)
        self.assertEqual(cemento['obras'], 2)

    def test_desglosa_por_obra(self):
        respuesta = self.client.get(reverse('pedidos:compras'))
        desglose = next(o for t, o in respuesta.context['totales']
                        if t['material_id'] == self.cemento.pk)
        cantidades = {o['nombre']: o['cantidad'] for o in desglose}
        self.assertEqual(cantidades, {'Obra Test': Decimal('30'), 'Obra Norte': Decimal('15')})

    def test_excluye_notas_cerradas(self):
        self.nota_beto.estado = 'cerrado'
        self.nota_beto.save(update_fields=['estado'])
        respuesta = self.client.get(reverse('pedidos:compras'))
        self.assertEqual(self.fila_de(respuesta, self.cemento)['total'], Decimal('30.00'))

    def test_excluye_notas_ya_pedidas(self):
        """Lo que ya está en una orden vigente no se vuelve a proponer."""
        orden(self.proveedor, notas=[self.nota_beto])
        respuesta = self.client.get(reverse('pedidos:compras'))
        self.assertEqual(self.fila_de(respuesta, self.cemento)['total'], Decimal('30.00'))

    def test_incluye_ya_pedidas_si_se_pide(self):
        orden(self.proveedor, notas=[self.nota_beto])
        respuesta = self.client.get(reverse('pedidos:compras'), {'incluir_pedidas': '1'})
        self.assertEqual(self.fila_de(respuesta, self.cemento)['total'], Decimal('45.00'))

    def test_una_orden_anulada_libera_sus_notas(self):
        orden(self.proveedor, notas=[self.nota_beto], anulada=True)
        respuesta = self.client.get(reverse('pedidos:compras'))
        self.assertEqual(self.fila_de(respuesta, self.cemento)['total'], Decimal('45.00'))

    def test_filtra_por_obra(self):
        obra_beto = self.nota_beto.obra
        respuesta = self.client.get(reverse('pedidos:compras'), {'obra': obra_beto.pk})
        cemento = self.fila_de(respuesta, self.cemento)
        self.assertEqual(cemento['total'], Decimal('15.00'))
        self.assertEqual(cemento['obras'], 1)

    def test_filtra_por_estado(self):
        self.nota_beto.estado = 'aprobada'
        self.nota_beto.save(update_fields=['estado'])
        respuesta = self.client.get(reverse('pedidos:compras'), {'estado': 'aprobada'})
        cemento = self.fila_de(respuesta, self.cemento)
        self.assertEqual(cemento['total'], Decimal('15.00'))

        respuesta = self.client.get(reverse('pedidos:compras'), {'estado': 'pendiente'})
        self.assertEqual(self.fila_de(respuesta, self.cemento)['total'], Decimal('30.00'))

    def test_filtra_por_rango_de_fechas(self):
        self.nota_ana.fecha = '2026-01-10'
        self.nota_ana.save(update_fields=['fecha'])
        self.nota_beto.fecha = '2026-09-10'
        self.nota_beto.save(update_fields=['fecha'])
        respuesta = self.client.get(reverse('pedidos:compras'), {'desde': '2026-06-01'})
        cemento = self.fila_de(respuesta, self.cemento)
        self.assertEqual(cemento['total'], Decimal('15.00'))

    def test_busqueda_por_texto(self):
        """'q' tiene que filtrar de verdad, no solo mostrarse en el input."""
        respuesta = self.client.get(reverse('pedidos:compras'), {'q': 'ladri'})
        self.assertIsNotNone(self.fila_de(respuesta, self.ladrillo))
        self.assertIsNone(self.fila_de(respuesta, self.cemento))

    def test_busqueda_por_nombre_del_material(self):
        respuesta = self.client.get(reverse('pedidos:compras'), {'q': 'CEMENTO'})
        self.assertIsNotNone(self.fila_de(respuesta, self.cemento))
        self.assertIsNone(self.fila_de(respuesta, self.ladrillo))

    def test_ignora_filtros_basura(self):
        for params in ({'obra': 'abc'}, {'estado': 'inventado'},
                       {'proveedor': 'abc'}, {'incluir_pedidas': 'si'}):
            with self.subTest(params=params):
                self.assertEqual(
                    self.client.get(reverse('pedidos:compras'), params).status_code, 200
                )

    def test_avisa_los_items_sin_catalogar(self):
        item(self.nota_ana, 'Something raro', 4)
        respuesta = self.client.get(reverse('pedidos:compras'), {'sin_catalogar': '1'})
        sin_catalogar = {f['detalle'] for f in respuesta.context['sin_catalogar']}
        self.assertIn('Something raro', sin_catalogar)

    def test_sin_catalogar_no_entra_al_grupo_de_totales(self):
        item(self.nota_ana, 'Something raro', 4)
        respuesta = self.client.get(reverse('pedidos:compras'), {'sin_catalogar': '1'})
        self.assertIsNone(self.fila_de(respuesta, material('Another')))


class OrdenCompraTest(BaseComprasTest):
    def test_crea_la_orden_con_cantidades_recalculadas(self):
        """Las cantidades las recalcula el servidor; el POST no las manda."""
        respuesta = self.crear_orden([self.cemento])
        self.assertEqual(respuesta.status_code, 302)
        creada = OrdenCompra.objects.get()
        self.assertEqual(creada.proveedor, self.proveedor)
        self.assertEqual(creada.lineas.get().cantidad, Decimal('45.00'))
        self.assertEqual(set(creada.notas.values_list('pk', flat=True)),
                         {self.nota_ana.pk, self.nota_beto.pk})
        self.assertEqual(creada.creado_por, self.admin)

    def test_guarda_el_precio_por_linea(self):
        self.crear_orden([self.cemento], **{f'precio_{self.cemento.pk}': '12,50'})
        creada = OrdenCompra.objects.get()
        self.assertEqual(creada.lineas.get().precio, Decimal('12.50'))

    def test_acepta_precio_con_punto(self):
        self.crear_orden([self.cemento], **{f'precio_{self.cemento.pk}': '9.75'})
        self.assertEqual(OrdenCompra.objects.get().lineas.get().precio, Decimal('9.75'))

    def test_precio_invalido_queda_en_cero(self):
        self.crear_orden([self.cemento], **{f'precio_{self.cemento.pk}': 'gratis'})
        self.assertEqual(OrdenCompra.objects.get().lineas.get().precio, Decimal('0.00'))

    def test_precio_negativo_se_ignora(self):
        self.crear_orden([self.cemento], **{f'precio_{self.cemento.pk}': '-5'})
        self.assertEqual(OrdenCompra.objects.get().lineas.get().precio, Decimal('0.00'))

    def test_el_proveedor_de_la_orden_no_filtra_las_notas(self):
        """Regresión: el filtro y la orden comparten el campo 'proveedor'.

        Las notas sin proveedor asignado (lo normal: lo elige el admin al
        comprar) no pueden quedar afuera por confundir el proveedor del filtro
        con el de la orden nueva.
        """
        self.nota_ana.proveedor = None
        self.nota_ana.save(update_fields=['proveedor'])
        self.nota_beto.proveedor = None
        self.nota_beto.save(update_fields=['proveedor'])
        # Solo el proveedor de la orden nueva: si se colara al filtro, la
        # consulta devolvería cero notas y la orden no se crearía.
        self.assertEqual(self.proveedor.pk, Proveedor.objects.first().pk)
        self._post_orden(
            materiales=[str(self.cemento.pk)],
            orden_proveedor=str(self.proveedor.pk),
        )
        creada = OrdenCompra.objects.get()
        self.assertEqual(creada.proveedor, self.proveedor)
        self.assertEqual(creada.lineas.get().cantidad, Decimal('45.00'))

    def test_respeta_el_filtro_de_obra_al_crear(self):
        """La orden se arma con el mismo recorte que se está mirando."""
        self.crear_orden([self.cemento], obra=str(self.nota_beto.obra.pk))
        creada = OrdenCompra.objects.get()
        self.assertEqual(creada.lineas.get().cantidad, Decimal('15.00'))
        self.assertEqual([n.pk for n in creada.notas.all()], [self.nota_beto.pk])

    def test_no_crea_orden_sin_materiales(self):
        respuesta = self.client.post(reverse('pedidos:compras_crear_orden'), {
            'orden_proveedor': str(self.proveedor.pk), 'orden_fecha': '2026-08-05',
        })
        self.assertEqual(respuesta.status_code, 302)
        self.assertEqual(OrdenCompra.objects.count(), 0)

    def test_exige_proveedor_y_fecha(self):
        respuesta = self.crear_orden([self.cemento], orden_proveedor='', orden_fecha='')
        self.assertEqual(respuesta.status_code, 302)
        self.assertEqual(OrdenCompra.objects.count(), 0)

    def test_ignora_materiales_inventados(self):
        self._post_orden(materiales=['999999'])
        self.assertEqual(OrdenCompra.objects.count(), 0)

    def test_no_crea_orden_si_no_queda_nada_por_comprar(self):
        orden(self.proveedor, notas=[self.nota_ana, self.nota_beto])
        self.crear_orden([self.cemento])
        self.assertEqual(OrdenCompra.objects.count(), 1)  # solo la anterior

    def test_marcar_cerrado_al_crear(self):
        self.crear_orden([self.cemento], marcar_cerrado='1')
        for nota in (self.nota_ana, self.nota_beto):
            nota.refresh_from_db()
            self.assertEqual(nota.estado, 'cerrado')
            self.assertEqual(nota.cerrado_por, 'Jefe Compras')

    def test_no_cierra_sin_marcarlo(self):
        self.crear_orden([self.cemento])
        self.nota_ana.refresh_from_db()
        self.assertEqual(self.nota_ana.estado, 'pendiente')

    def test_numeracion_consecutiva(self):
        """Las dos órdenes necesitan notas distintas: la 1ª consume las suyas."""
        self.crear_orden([self.cemento])
        item(self.nota_huerfana, 'Ladrillo hueco', 40)
        self.crear_orden([self.ladrillo])
        self.assertEqual(sorted(o.numero for o in OrdenCompra.objects.all()), [1, 2])

    def test_get_no_crea_nada(self):
        self.assertEqual(
            self.client.get(reverse('pedidos:compras_crear_orden')).status_code, 302
        )
        self.assertEqual(OrdenCompra.objects.count(), 0)

    def test_un_solo_material_no_pide_todos_los_montos(self):
        """Elegir solo ladrillo no debe colarse el cemento en la orden."""
        self.crear_orden([self.ladrillo])
        creada = OrdenCompra.objects.get()
        self.assertEqual([l.material for l in creada.lineas.all()], [self.ladrillo])
        self.assertEqual(creada.lineas.get().cantidad, Decimal('100.00'))


class AnulacionOrdenTest(BaseComprasTest):
    def test_anular_devuelve_las_notas_al_consolidado(self):
        self.crear_orden([self.cemento])
        creada = OrdenCompra.objects.get()
        respuesta = self.client.get(reverse('pedidos:compras'))
        self.assertIsNone(self.fila_de(respuesta, self.cemento))

        self.client.post(reverse('pedidos:orden_anular', args=[creada.pk]))
        creada.refresh_from_db()
        self.assertTrue(creada.anulada)

        respuesta = self.client.get(reverse('pedidos:compras'))
        self.assertEqual(self.fila_de(respuesta, self.cemento)['total'], Decimal('45.00'))

    def test_anular_dos_veces_no_rompe(self):
        self.crear_orden([self.cemento])
        creada = OrdenCompra.objects.get()
        url = reverse('pedidos:orden_anular', args=[creada.pk])
        self.assertEqual(self.client.post(url).status_code, 302)
        self.assertEqual(self.client.post(url).status_code, 302)
        creada.refresh_from_db()
        self.assertTrue(creada.anulada)

    def test_anular_reabre_lo_que_la_orden_cerio(self):
        """Si la orden cerró las notas, anularlas devuelve la demanda al consolidado."""
        self.crear_orden([self.cemento], marcar_cerrado='1')
        creada = OrdenCompra.objects.get()
        for nota in creada.notas.all():
            self.assertEqual(nota.estado, 'cerrado')

        respuesta = self.client.get(reverse('pedidos:compras'))
        self.assertIsNone(self.fila_de(respuesta, self.cemento))

        self.client.post(reverse('pedidos:orden_anular', args=[creada.pk]))
        self.assertEqual(self.fila_de(self.client.get(reverse('pedidos:compras')),
                                     self.cemento)['total'], Decimal('45.00'))
        for nota in creada.notas.all():
            nota.refresh_from_db()
            self.assertEqual(nota.estado, 'pendiente')
            self.assertEqual(nota.cerrado_por, '')

    def test_anular_respeta_una_que_ya_estaba_cerrada(self):
        """Una nota cerrada antes de la orden no se reabre al anularla."""
        self.nota_beto.estado = 'cerrado'
        self.nota_beto.save(update_fields=['estado'])
        self.crear_orden([self.cemento], marcar_cerrado='1')
        creada = OrdenCompra.objects.get()

        self.client.post(reverse('pedidos:orden_anular', args=[creada.pk]))
        self.nota_beto.refresh_from_db()
        self.assertEqual(self.nota_beto.estado, 'cerrado')
        self.nota_ana.refresh_from_db()
        self.assertEqual(self.nota_ana.estado, 'pendiente')

    def test_anular_no_toca_notas_de_otras_ordenes(self):
        primera = OrdenCompra.objects.create(
            numero=1, proveedor=self.proveedor, fecha='2026-08-01')
        primera.notas.set([self.nota_beto])
        self.nota_beto.estado = 'cerrado'
        self.nota_beto.save(update_fields=['estado'])

        segunda = OrdenCompra.objects.create(
            numero=2, proveedor=self.proveedor, fecha='2026-08-02')
        segunda.notas.set([self.nota_ana])

        self.client.post(reverse('pedidos:orden_anular', args=[segunda.pk]))
        self.nota_beto.refresh_from_db()
        self.assertEqual(self.nota_beto.estado, 'cerrado', 'pertenece a otra orden')

    def test_get_no_anula(self):
        self.crear_orden([self.cemento])
        creada = OrdenCompra.objects.get()
        self.client.get(reverse('pedidos:orden_anular', args=[creada.pk]))
        creada.refresh_from_db()
        self.assertFalse(creada.anulada)

    def test_totales_de_la_orden(self):
        self.crear_orden([self.cemento], **{f'precio_{self.cemento.pk}': '10'})
        creada = OrdenCompra.objects.get()
        # 45 unidades x 10
        self.assertEqual(creada.total, Decimal('450.00'))


class ExcelOrdenTest(BaseComprasTest):
    def test_descarga_excel_de_la_orden(self):
        self.crear_orden([self.cemento], **{f'precio_{self.cemento.pk}': '10'})
        creada = OrdenCompra.objects.get()
        respuesta = self.client.get(reverse('pedidos:orden_descargar', args=[creada.pk]))
        self.assertEqual(respuesta.status_code, 200)
        self.assertIn('spreadsheetml', respuesta['Content-Type'])
        self.assertIn(f'OC_{creada.numero:04d}', respuesta['Content-Disposition'])

    def test_excel_incluye_precio_y_total(self):
        from openpyxl import load_workbook
        import io

        self.crear_orden([self.cemento], **{f'precio_{self.cemento.pk}': '10'})
        creada = OrdenCompra.objects.get()
        respuesta = self.client.get(reverse('pedidos:orden_descargar', args=[creada.pk]))
        hoja = load_workbook(io.BytesIO(respuesta.content)).active
        # Fila 10 = primer material (encabezado en 9)
        self.assertEqual(hoja.cell(row=10, column=5).value, 10.0)
        self.assertEqual(hoja.cell(row=10, column=6).value, 450.0)
        # Fila 11 = total
        self.assertEqual(hoja['E11'].value, 'TOTAL')
        self.assertEqual(hoja.cell(row=11, column=6).value, 450.0)

    def test_excel_fecha_en_formato_dia_mes_anio(self):
        """La fecha va como valor real de Excel con formato dd/mm/yyyy."""
        from openpyxl import load_workbook
        import io

        creada = orden()
        respuesta = self.client.get(reverse('pedidos:orden_descargar', args=[creada.pk]))
        hoja = load_workbook(io.BytesIO(respuesta.content)).active
        celda = hoja['E5']
        self.assertEqual(celda.number_format, 'dd/mm/yyyy')
        # Como datetime, no como texto: así Excel lo ordena y lo usa en fórmulas.
        self.assertEqual(celda.value.date(), date(2026, 8, 1))

    def test_excel_lista_las_notas_de_origen(self):
        from openpyxl import load_workbook
        import io

        self.crear_orden([self.cemento])
        creada = OrdenCompra.objects.get()
        respuesta = self.client.get(reverse('pedidos:orden_descargar', args=[creada.pk]))
        hoja = load_workbook(io.BytesIO(respuesta.content)).active
        textos = [str(c.value) for fila in hoja.iter_rows() for c in fila if c.value]
        self.assertTrue(any(t.startswith('NP ') for t in textos), textos)


class PermisosComprasTest(BasePermisosTest):
    def setUp(self):
        self.cemento = material('Cemento')
        item(self.nota_ana, 'Cemento', 10)

    def test_solicitante_no_entra_a_compras(self):
        self.client.force_login(self.ana)
        self.assertEqual(self.client.get(reverse('pedidos:compras')).status_code, 403)

    def test_solicitante_no_crea_orden(self):
        self.client.force_login(self.ana)
        respuesta = self.client.post(reverse('pedidos:compras_crear_orden'), {
            'materiales': [str(self.cemento.pk)],
        })
        self.assertEqual(respuesta.status_code, 403)
        self.assertEqual(OrdenCompra.objects.count(), 0)

    def test_solicitante_no_ve_ordenes(self):
        creada = orden(notas=[self.nota_ana])
        self.client.force_login(self.ana)
        self.assertEqual(
            self.client.get(reverse('pedidos:orden_detalle', args=[creada.pk])).status_code, 403
        )

    def test_solicitante_no_descarga_excel(self):
        creada = orden(notas=[self.nota_ana])
        self.client.force_login(self.ana)
        self.assertEqual(
            self.client.get(reverse('pedidos:orden_descargar', args=[creada.pk])).status_code, 403
        )

    def test_solicitante_no_anula(self):
        creada = orden(notas=[self.nota_ana])
        self.client.force_login(self.ana)
        self.assertEqual(
            self.client.post(reverse('pedidos:orden_anular', args=[creada.pk])).status_code, 403
        )
        creada.refresh_from_db()
        self.assertFalse(creada.anulada)

    def test_requiere_login(self):
        self.client.logout()
        self.assertEqual(self.client.get(reverse('pedidos:compras')).status_code, 302)