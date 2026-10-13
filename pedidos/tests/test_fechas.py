"""La exportación a Excel debe usar d/m/a en las fechas."""

import io
from datetime import date

from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from pedidos.excel import EXCEL_DATE_FORMAT, _fmt_fecha
from pedidos.models import Material, NotaPedido, Obra, OrdenCompra, Proveedor

from .base import linea, orden


class FormatoFechaTest(TestCase):
    def test_formato_declarado_es_dd_mm_yyyy(self):
        self.assertEqual(EXCEL_DATE_FORMAT, 'dd/mm/yyyy')

    def test_fmt_fecha_devuelve_fecha_real_no_texto(self):
        """Un ISO en texto no lo reconoce Excel como fecha."""
        f = date(2026, 3, 15)
        self.assertEqual(_fmt_fecha(f), f)
        self.assertNotIsInstance(_fmt_fecha(f), str)

    def test_fmt_fecha_sin_fecha_devuelve_none(self):
        self.assertIsNone(_fmt_fecha(None))


class ExcelNotaFechaTest(TestCase):
    """Las celdas de fecha de la plantilla (K14 fecha, K15 límite)."""

    def setUp(self):
        self.obra = Obra.objects.create(nombre='Obra Fechas')
        self.nota = NotaPedido.objects.create(
            numero=1, obra=self.obra, fecha=date(2026, 3, 15),
            limite=date(2026, 4, 20), solicitante='Ana',
        )
        from django.contrib.auth.models import User
        user = User.objects.create_user('admin', password='clave12345', is_superuser=True)
        self.client.force_login(user)

    def _hoja(self):
        respuesta = self.client.get(reverse('pedidos:descargar', args=[self.nota.pk]))
        self.assertEqual(respuesta.status_code, 200)
        return load_workbook(io.BytesIO(respuesta.content)).active

    def test_fecha_de_la_nota(self):
        celda = self._hoja()['K14']
        self.assertEqual(celda.number_format, EXCEL_DATE_FORMAT)
        self.assertEqual(celda.value.date(), date(2026, 3, 15))

    def test_fecha_limite(self):
        celda = self._hoja()['K15']
        self.assertEqual(celda.number_format, EXCEL_DATE_FORMAT)
        self.assertEqual(celda.value.date(), date(2026, 4, 20))

    def test_sin_limite_no_rompe(self):
        self.nota.limite = None
        self.nota.save(update_fields=['limite'])
        celda = self._hoja()['K15']
        self.assertIsNone(celda.value)


class ExcelOrdenFechaTest(TestCase):
    def setUp(self):
        from django.contrib.auth.models import User
        user = User.objects.create_user('admin', password='clave12345', is_superuser=True)
        self.client.force_login(user)
        self.proveedor = Proveedor.objects.create(nombre='Prov Fechas')
        self.material = Material.objects.create(nombre='Cemento')
        self.orden = orden(self.proveedor, fecha=date(2026, 12, 31))
        linea(self.orden, self.material, 2, precio=10)

    def test_fecha_de_la_orden(self):
        respuesta = self.client.get(
            reverse('pedidos:orden_descargar', args=[self.orden.pk])
        )
        hoja = load_workbook(io.BytesIO(respuesta.content)).active
        celda = hoja['E5']
        self.assertEqual(celda.number_format, EXCEL_DATE_FORMAT)
        self.assertEqual(celda.value.date(), date(2026, 12, 31))

    def test_subtotal_y_total_siguen_siendo_numeros(self):
        """Sumar no debe romperse por el cambio de formato de fecha."""
        respuesta = self.client.get(
            reverse('pedidos:orden_descargar', args=[self.orden.pk])
        )
        hoja = load_workbook(io.BytesIO(respuesta.content)).active
        self.assertEqual(hoja.cell(row=10, column=6).value, 20.0)
        self.assertEqual(hoja.cell(row=11, column=6).value, 20.0)